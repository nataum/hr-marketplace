#!/usr/bin/env python3
"""Construye results.xlsx directamente desde datos estructurados (sin results.md).

Reutiliza EXACTAMENTE el formato del job-matcher original (export_xlsx.py):
mismas columnas, mismo header azul, bordes, columnas Rationale/URL justificadas,
resaltado de Fit (Low=rojo, High=amarillo), hipervínculos en URL, dropdown de
Application Status, ancho de Notes, fila header congelada y autofiltro. Dedup por
URL y, si se pasa --previous <xlsx>, preserva Application Status / Notes por fila.

Entrada JSON (--data data.json):
{
  "matches": [
    {"Company": "...", "Role": "...", "Fit": "High|Medium|Low",
     "Rationale": "...", "Date Published": "...", "Detection Date": "YYYY-MM-DD",
     "Location/Modality": "...", "Salary": "...", "URL": "..."},
    ...
  ],
  "no_match": [
    {"Company": "...", "Reason": "..."},
    ...
  ]
}

Uso:
    python3 scripts/build_xlsx.py --data data.json --out /mnt/user-data/outputs/results.xlsx [--previous prev.xlsx]
"""
import argparse
import json
import os

import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    "Company", "Role", "Fit", "Rationale", "Date Published", "Detection Date",
    "Location/Modality", "Salary", "URL", "Application Status", "Notes",
]

FIT_FILLS = {
    "low": PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid"),
    "high": PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid"),
}

APPLICATION_STATUS_OPTIONS = ["Applied", "Discarded", "I'm not interested"]
NOTES_MIN_WIDTH = 30


def build_rows(data):
    rows = []
    for m in data.get("matches", []):
        rows.append({
            "Company": (m.get("Company") or "").strip(),
            "Role": (m.get("Role") or "").strip(),
            "Fit": (m.get("Fit") or "").strip(),
            "Rationale": (m.get("Rationale") or "").strip(),
            "Date Published": (m.get("Date Published") or "").strip(),
            "Detection Date": (m.get("Detection Date") or "").strip(),
            "Location/Modality": (m.get("Location/Modality") or "").strip(),
            "Salary": (m.get("Salary") or "").strip(),
            "URL": (m.get("URL") or "").strip(),
            "Application Status": "Not reviewed",
            "Notes": "",
        })
    for nm in data.get("no_match", []):
        rows.append({
            "Company": (nm.get("Company") or "").strip(),
            "Role": "",
            "Fit": "No Match",
            "Rationale": " ".join((nm.get("Reason") or "").split()),
            "Date Published": "",
            "Detection Date": "",
            "Location/Modality": "",
            "Salary": "",
            "URL": "",
            "Application Status": "",
            "Notes": "",
        })
    return rows


def dedupe_by_url(rows):
    seen, deduped, dropped = set(), [], 0
    for row in rows:
        url = row["URL"]
        if url and url in seen:
            dropped += 1
            continue
        if url:
            seen.add(url)
        deduped.append(row)
    return deduped, dropped


def row_key(d):
    return d["URL"] if d["URL"] else f'{d["Company"]}::{d["Role"]}'


def load_previous(path):
    if not path or not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if "Company" not in header:
        return {}
    idx = {name: header.index(name) for name in header if name}
    old = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {name: (r[i] if r[i] is not None else "") for name, i in idx.items()}
        old[row_key(d)] = d
    return old


def preserve_tracking_fields(rows, old_by_key):
    for row in rows:
        old = old_by_key.get(row_key(row))
        if old:
            if old.get("Application Status"):
                row["Application Status"] = old["Application Status"]
            if old.get("Notes"):
                row["Notes"] = old["Notes"]


def apply_formatting(out_path):
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    header_fill = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFFFF")
    hyperlink_font = Font(color="FF0563C1", underline="single")
    thin = Side(style="thin", color="FF000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    max_row = ws.max_row
    max_col = ws.max_column

    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_letters = {ws.cell(row=1, column=i).value: i for i in range(1, max_col + 1)}
    justify_cols = {col_letters.get("URL"), col_letters.get("Rationale")}
    fit_col = col_letters.get("Fit")
    url_col = col_letters.get("URL")
    status_col = col_letters.get("Application Status")

    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            if row_idx > 1 and col_idx in justify_cols:
                cell.alignment = Alignment(horizontal="justify", vertical="top", wrap_text=True)
            elif row_idx > 1:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    if fit_col:
        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=fit_col)
            fit_value = str(cell.value or "").strip().lower()
            fill = FIT_FILLS.get(fit_value)
            if fill:
                cell.fill = fill

    if url_col:
        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=url_col)
            url = str(cell.value or "").strip()
            if url.startswith("http"):
                cell.hyperlink = url
                cell.font = hyperlink_font

    if status_col:
        status_letter = get_column_letter(status_col)
        options = ",".join(APPLICATION_STATUS_OPTIONS)
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{status_letter}2:{status_letter}{max_row}")

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        header_name = ws.cell(row=1, column=col_idx).value
        values = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, max_row + 1)]
        width = min(max(len(v) for v in values) + 2, 60)
        if header_name == "Notes":
            width = max(width, NOTES_MIN_WIDTH)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    wb.save(out_path)


def write_workbook(rows, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", required=True, help="JSON con matches y no_match")
    parser.add_argument("--out", default="/mnt/user-data/outputs/results.xlsx")
    parser.add_argument("--previous", default=None, help="Excel de una corrida anterior (opcional) para preservar Application Status/Notes")
    args = parser.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    rows = build_rows(data)
    rows, dropped = dedupe_by_url(rows)

    old_by_key = load_previous(args.previous)
    preserve_tracking_fields(rows, old_by_key)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    write_workbook(rows, args.out)
    apply_formatting(args.out)

    print(f"Wrote {len(rows)} rows to {args.out} ({dropped} duplicate URL row(s) dropped)")
    if old_by_key:
        print(f"Preserved tracking fields from previous workbook ({len(old_by_key)} prior rows).")


if __name__ == "__main__":
    main()
